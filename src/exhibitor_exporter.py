"""
PDF Exhibitor Data Exporter — Excel Writer
==========================================
Author  : Ashwin
GitHub  : https://github.com/ashwin8332
Version : 3.0.0

Formats extracted exhibitor dicts into a styled, two-sheet Excel workbook:

  Sheet 1 — "Exhibitor Data"
      S.No. | Company Name | Address | Contact Person |
      Tel./Mobile | E-mail | Products on Display

      Style features:
        - Dark-navy header row with white bold text
        - Alternating light-blue row striping
        - Auto-filter on all columns
        - Frozen header row (first row stays visible while scrolling)
        - Wrapped text in multi-line cells
        - Title row with PDF source name and author credit

  Sheet 2 — "Summary"
      Key statistics: total exhibitors, field-level completeness percentages.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ── Column definitions ─────────────────────────────────────────────────────────

# Maps internal extractor keys -> human-readable Excel column headers
COLUMN_MAP: dict[str, str] = {
    "company_name":        "Company Name",
    "address":             "Address",
    "contact_person":      "Contact Person",
    "tel_mobile":          "Tel./Mobile",
    "email":               "E-mail",
    "products_on_display": "Products on Display",
}

# Excel column widths (characters)
COLUMN_WIDTHS: dict[str, int] = {
    "S.No.":               6,
    "Company Name":        35,
    "Address":             45,
    "Contact Person":      22,
    "Tel./Mobile":         20,
    "E-mail":              32,
    "Products on Display": 55,
}

# ── Styles ─────────────────────────────────────────────────────────────────────
_HDR_FILL    = PatternFill("solid", fgColor="1F3864")
_HDR_FONT    = Font(name="Calibri", bold=True,  color="FFFFFF", size=11)
_TITLE_FONT  = Font(name="Calibri", bold=True,  color="1F3864", size=14)
_BODY_FONT   = Font(name="Calibri",              size=10)
_ALT_FILL    = PatternFill("solid", fgColor="EAF0FB")
_THIN_BORDER = Border(
    left=Side(style="thin",   color="CCCCCC"),
    right=Side(style="thin",  color="CCCCCC"),
    top=Side(style="thin",    color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
_CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_W  = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
_LEFT_C  = Alignment(horizontal="left",   vertical="center", wrap_text=False)


# ── Public API ─────────────────────────────────────────────────────────────────

def export_to_excel(
    exhibitors: list[dict[str, str]],
    output_path: str | Path | None = None,
    pdf_filename: str = "unknown.pdf",
) -> str | BytesIO:
    """
    Export *exhibitors* to a formatted Excel workbook.

    Parameters
    ----------
    exhibitors  : list of dicts from ``exhibitor_extractor.extract_exhibitors``
    output_path : write to disk and return path str, or return BytesIO if None
    pdf_filename: displayed in the workbook header/summary

    Returns
    -------
    str (file path) or BytesIO buffer
    """
    df = _build_df(exhibitors)
    wb = openpyxl.Workbook()
    _write_data_sheet(wb, df, pdf_filename)
    _write_summary_sheet(wb, df, pdf_filename)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    if output_path is not None:
        Path(output_path).write_bytes(buf.read())
        return str(output_path)

    buf.seek(0)
    return buf


# ── DataFrame builder ──────────────────────────────────────────────────────────

def _build_df(exhibitors: list[dict[str, str]]) -> pd.DataFrame:
    rows = [
        {display: ex.get(internal, "")
         for internal, display in COLUMN_MAP.items()}
        for ex in exhibitors
    ]
    df = pd.DataFrame(rows, columns=list(COLUMN_MAP.values()))
    df.insert(0, "S.No.", range(1, len(df) + 1))
    return df


# ── Data sheet ─────────────────────────────────────────────────────────────────

def _write_data_sheet(wb: openpyxl.Workbook, df: pd.DataFrame, pdf_filename: str) -> None:
    ws = wb.active
    ws.title = "Exhibitor Data"
    ncols = len(df.columns)

    # Row 1 — Title
    ws.append([f"Exhibitor Directory — {pdf_filename}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    tc = ws.cell(1, 1)
    tc.font = _TITLE_FONT
    tc.alignment = _CENTER
    ws.row_dimensions[1].height = 28

    # Row 2 — Sub-header
    ws.append([
        f"Extracted by: Ashwin  |  "
        f"GitHub: https://github.com/ashwin8332  |  "
        f"Total: {len(df)} exhibitors"
    ])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    sc = ws.cell(2, 1)
    sc.font = Font(name="Calibri", italic=True, color="555555", size=9)
    sc.alignment = _CENTER
    ws.row_dimensions[2].height = 16

    # Row 3 — blank spacer
    ws.append([])
    ws.row_dimensions[3].height = 6

    # Row 4 — column headers
    ws.append(list(df.columns))
    for col_idx, col_name in enumerate(df.columns, start=1):
        c = ws.cell(4, col_idx)
        c.font = _HDR_FONT
        c.fill = _HDR_FILL
        c.alignment = _CENTER
        c.border = _THIN_BORDER
    ws.row_dimensions[4].height = 22

    # Data rows (start at row 5)
    for row_num, row_data in enumerate(df.itertuples(index=False), start=5):
        ws.append(list(row_data))
        even = (row_num - 4) % 2 == 0
        for col_idx in range(1, ncols + 1):
            c = ws.cell(row_num, col_idx)
            c.font   = _BODY_FONT
            c.border = _THIN_BORDER
            c.alignment = _LEFT_W if col_idx > 2 else _LEFT_C
            if even:
                c.fill = _ALT_FILL

    # Column widths
    for col_idx, col_name in enumerate(df.columns, start=1):
        ltr = get_column_letter(col_idx)
        ws.column_dimensions[ltr].width = COLUMN_WIDTHS.get(col_name, 20)

    ws.freeze_panes = "A5"
    last_col = get_column_letter(ncols)
    ws.auto_filter.ref = f"A4:{last_col}{4 + len(df)}"


# ── Summary sheet ──────────────────────────────────────────────────────────────

def _write_summary_sheet(wb: openpyxl.Workbook, df: pd.DataFrame, pdf_filename: str) -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 45

    def _row(label: str, value: str, bold: bool = False) -> None:
        ws.append([label, value])
        r = ws.max_row
        ws.cell(r, 1).font = Font(name="Calibri", bold=bold, size=10)
        ws.cell(r, 2).font = Font(name="Calibri", size=10)
        ws.cell(r, 1).alignment = _LEFT_C
        ws.cell(r, 2).alignment = _LEFT_C

    # Title
    ws.append(["Exhibitor Extraction Summary"])
    ws.merge_cells("A1:B1")
    tc = ws.cell(1, 1)
    tc.font = _TITLE_FONT
    tc.alignment = _CENTER
    ws.row_dimensions[1].height = 26
    ws.append([])

    _row("Source PDF",       pdf_filename,                         bold=True)
    _row("Total Exhibitors", str(len(df)),                         bold=True)
    _row("Author",           "Ashwin",                             bold=True)
    _row("GitHub",           "https://github.com/ashwin8332",      bold=True)
    ws.append([])

    # Field completeness header
    ws.append(["Field Completeness"])
    ws.merge_cells(f"A{ws.max_row}:B{ws.max_row}")
    hc = ws.cell(ws.max_row, 1)
    hc.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    hc.fill = _HDR_FILL
    hc.alignment = _CENTER

    data_cols = [c for c in df.columns if c != "S.No."]
    for col in data_cols:
        filled = df[col].apply(lambda v: bool(str(v).strip())).sum()
        pct    = round(filled / len(df) * 100, 1) if len(df) else 0
        _row(col, f"{filled} / {len(df)} ({pct}%)")
